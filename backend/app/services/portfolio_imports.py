"""Bounded file inspection, mapping, validation, and persistence for imports."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import PurePosixPath
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from backend.app.core.settings import get_settings
from backend.app.models.company import Company, CompanyFinancials, CreditProfile, PortfolioExposure
from backend.app.models.portfolio_import import PortfolioImport

MAX_XLSX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_XLSX_MEMBERS = 250
MAX_XLSX_COMPRESSION_RATIO = 100

CANONICAL_FIELDS: dict[str, dict[str, str]] = {
    "company_name": {"label": "Company name", "required": "true"},
    "external_id": {"label": "External company ID", "required": "false"},
    "industry": {"label": "Industry", "required": "false"},
    "country": {"label": "Country", "required": "false"},
    "currency": {"label": "Reporting currency (ISO 4217)", "required": "conditional"},
    "fiscal_year": {"label": "Financial reporting year", "required": "conditional"},
    "annual_revenue": {"label": "Annual revenue", "required": "false"},
    "revenue": {"label": "Statement revenue", "required": "false"},
    "ebitda": {"label": "EBITDA", "required": "false"},
    "net_income": {"label": "Net income", "required": "false"},
    "total_assets": {"label": "Total assets", "required": "false"},
    "total_liabilities": {"label": "Total liabilities", "required": "false"},
    "current_assets": {"label": "Current assets", "required": "false"},
    "current_liabilities": {"label": "Current liabilities", "required": "false"},
    "debt": {"label": "Total debt", "required": "false"},
    "cash": {"label": "Cash and equivalents", "required": "false"},
    "operating_cash_flow": {"label": "Operating cash flow", "required": "false"},
    "exposure": {"label": "Total exposure", "required": "false"},
    "outstanding_amount": {"label": "Outstanding principal", "required": "false"},
    "requested_loan_amount": {"label": "Requested loan amount", "required": "false"},
    "interest_rate": {"label": "Interest rate (percent)", "required": "false"},
    "delinquency_days": {"label": "Days past due", "required": "false"},
}

ALIASES: dict[str, set[str]] = {
    "company_name": {"company", "companyname", "borrower", "borrowername", "businessname", "entityname", "name"},
    "external_id": {"externalid", "companyid", "borrowerid", "customerid", "entityid", "reference"},
    "industry": {"industry", "sector", "industrysector"},
    "country": {"country", "countrycode"},
    "currency": {"currency", "currencycode", "currencyiso", "reportingcurrency"},
    "fiscal_year": {"fiscalyear", "financialyear", "reportingyear", "year"},
    "annual_revenue": {"annualrevenue", "revenueannual"},
    "revenue": {"revenue", "sales", "turnover", "totalrevenue"},
    "ebitda": {"ebitda"},
    "net_income": {"netincome", "profitaftertax", "pat"},
    "total_assets": {"totalassets", "assets"},
    "total_liabilities": {"totalliabilities", "liabilities"},
    "current_assets": {"currentassets"},
    "current_liabilities": {"currentliabilities"},
    "debt": {"totaldebt", "debt", "borrowings"},
    "cash": {"cash", "cashequivalents", "cashandcashequivalents"},
    "operating_cash_flow": {"operatingcashflow", "cashflowfromoperations"},
    "exposure": {"exposure", "totalexposure", "sanctionedlimit"},
    "outstanding_amount": {"outstanding", "outstandingprincipal", "principaloutstanding"},
    "requested_loan_amount": {"loanamount", "requestedloanamount", "sanctionedamount"},
    "interest_rate": {"interestrate", "rateofinterest"},
    "delinquency_days": {"dpd", "dayspastdue", "delinquencydays"},
}

FINANCIAL_FIELDS = {
    "revenue", "ebitda", "net_income", "total_assets", "total_liabilities",
    "current_assets", "current_liabilities", "debt", "cash", "operating_cash_flow",
}
NONNEGATIVE_FIELDS = {
    "annual_revenue", "revenue", "total_assets", "total_liabilities", "current_assets",
    "current_liabilities", "debt", "cash", "exposure", "outstanding_amount",
    "requested_loan_amount",
}
MONETARY_FIELDS = NONNEGATIVE_FIELDS | FINANCIAL_FIELDS | {"interest_rate"}
ISO_CURRENCIES = frozenset(
    "AED AFN ALL AMD ANG AOA ARS AUD AWG AZN BAM BBD BDT BGN BHD BIF BMD BND BOB BOV BRL BSD BTN BWP BYN BZD CAD CDF CHE CHF CHW CLF CLP CNY COP COU CRC CUC CUP CVE CZK DJF DKK DOP DZD EGP ERN ETB EUR FJD FKP GBP GEL GHS GIP GMD GNF GTQ GYD HKD HNL HTG HUF IDR ILS INR IQD IRR ISK JMD JOD JPY KES KGS KHR KMF KPW KRW KWD KYD KZT LAK LBP LKR LRD LSL LYD MAD MDL MGA MKD MMK MNT MOP MRU MUR MVR MWK MXN MXV MYR MZN NAD NGN NIO NOK NPR NZD OMR PAB PEN PGK PHP PKR PLN PYG QAR RON RSD RUB RWF SAR SBD SCR SDG SEK SGD SHP SLE SLL SOS SRD SSP STN SVC SYP SZL THB TJS TMT TND TOP TRY TTD TWD TZS UAH UGX USD USN UYI UYU UYW UZS VED VES VND VUV WST XAF XAG XAU XBA XBB XBC XBD XCD XCG XDR XOF XPD XPF XPT XSU XTS XUA XXX YER ZAR ZMW ZWL".split()
)


class ImportFileError(ValueError):
    """An upload is malformed or exceeds a configured safety limit."""


class ImportStateError(ValueError):
    """An import operation is invalid for the current state."""


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _headers_key(header: str) -> str:
    return re.sub(r"[^a-z0-9]", "", header.casefold())


def _validate_headers(headers: list[str]) -> None:
    settings = get_settings()
    if not headers or all(not h.strip() for h in headers):
        raise ImportFileError("The selected worksheet has no header row.")
    if len(headers) > settings.portfolio_import_max_columns:
        raise ImportFileError("The file exceeds the column limit.")
    keys = [_headers_key(h.strip()) for h in headers]
    if any(not key for key in keys):
        raise ImportFileError("Column headers cannot be blank.")
    if len(keys) != len(set(keys)):
        raise ImportFileError("Duplicate column headers are not supported.")


def suggest_mapping(headers: list[str]) -> dict[str, str | None]:
    suggestions: dict[str, str | None] = {}
    used: set[str] = set()
    for header in headers:
        normalized = _headers_key(header.strip())
        suggestion = next(
            (field for field, aliases in ALIASES.items() if normalized in aliases), None
        )
        if suggestion in used:
            suggestion = None
        if suggestion:
            used.add(suggestion)
        suggestions[header] = suggestion
    return suggestions


def _csv_sheets(content: bytes) -> dict[str, dict[str, Any]]:
    if b"\x00" in content:
        raise ImportFileError("The CSV file contains invalid binary data.")
    try:
        text = content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text, newline=""), strict=True)
        headers = next(reader, None)
        if headers is None:
            raise ImportFileError("The file is empty.")
        headers = [str(header).strip() for header in headers]
        _validate_headers(headers)
        rows: list[dict[str, Any]] = []
        max_rows = get_settings().portfolio_import_max_rows
        for row_number, values in enumerate(reader, start=2):
            if not values or all(not str(value).strip() for value in values):
                continue
            if len(values) != len(headers):
                raise ImportFileError(f"Row {row_number} has a different number of columns.")
            if len(rows) >= max_rows:
                raise ImportFileError("The file exceeds the row limit.")
            rows.append({
                "row_number": row_number,
                "values": [
                    {"value": value, "formula": value.lstrip().startswith("=")}
                    for value in values
                ],
            })
    except UnicodeDecodeError as exc:
        raise ImportFileError("CSV files must use UTF-8 encoding.") from exc
    except csv.Error as exc:
        raise ImportFileError("The CSV file is malformed.") from exc
    if not rows:
        raise ImportFileError("The file contains no data rows.")
    return {"CSV": {"headers": headers, "rows": rows}}


def _check_xlsx_archive(content: bytes) -> None:
    if not content.startswith(b"PK"):
        raise ImportFileError("The XLSX file content is invalid.")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_XLSX_MEMBERS:
                raise ImportFileError("The workbook contains too many internal files.")
            total_size = sum(member.file_size for member in members)
            if total_size > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ImportFileError("The workbook expands beyond the allowed size.")
            for member in members:
                if member.flag_bits & 0x1:
                    raise ImportFileError("Encrypted workbooks are not supported.")
                if member.file_size > 0 and member.compress_size == 0:
                    raise ImportFileError("The workbook contains an invalid compressed file.")
                if member.compress_size and member.file_size / member.compress_size > MAX_XLSX_COMPRESSION_RATIO:
                    raise ImportFileError("The workbook compression ratio exceeds the safety limit.")
                if ".." in PurePosixPath(member.filename).parts:
                    raise ImportFileError("The workbook contains an invalid internal path.")
    except zipfile.BadZipFile as exc:
        raise ImportFileError("The XLSX workbook is malformed.") from exc


def _xlsx_sheets(content: bytes) -> dict[str, dict[str, Any]]:
    _check_xlsx_archive(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        sheets: dict[str, dict[str, Any]] = {}
        remaining_rows = get_settings().portfolio_import_max_rows
        for worksheet in workbook.worksheets:
            if worksheet.sheet_state != "visible":
                continue
            header_cells = next(worksheet.iter_rows(max_row=1), None)
            if not header_cells:
                continue
            headers = [str(_json_value(cell.value) or "").strip() for cell in header_cells]
            while headers and not headers[-1]:
                headers.pop()
            _validate_headers(headers)
            rows: list[dict[str, Any]] = []
            for row_number, cells in enumerate(worksheet.iter_rows(min_row=2), start=2):
                cells = cells[:len(headers)]
                if not any(cell.value is not None for cell in cells):
                    continue
                if len(rows) >= remaining_rows:
                    raise ImportFileError("The workbook exceeds the row limit.")
                rows.append({
                    "row_number": row_number,
                    "values": [
                        {"value": _json_value(cell.value), "formula": cell.data_type == "f"}
                        for cell in cells
                    ],
                })
            remaining_rows -= len(rows)
            if rows:
                sheets[worksheet.title] = {"headers": headers, "rows": rows}
        workbook.close()
    except ImportFileError:
        raise
    except Exception as exc:
        raise ImportFileError("The XLSX workbook is malformed or unsupported.") from exc
    if not sheets:
        raise ImportFileError("The workbook has no visible worksheets with data.")
    return sheets


def inspect_upload(filename: str, content: bytes) -> tuple[str, dict[str, Any]]:
    """Inspect a bounded CSV/XLSX file without using its name as a path."""
    suffix = PurePosixPath(filename.replace("\\", "/")).suffix.casefold()
    if suffix not in {".csv", ".xlsx"}:
        raise ImportFileError("Only CSV and XLSX files are supported.")
    if not content:
        raise ImportFileError("The uploaded file is empty.")
    sheets = _csv_sheets(content) if suffix == ".csv" else _xlsx_sheets(content)
    return suffix[1:], sheets


def safe_filename(filename: str) -> str:
    name = PurePosixPath((filename or "upload").replace("\\", "/")).name
    name = re.sub(r"[^A-Za-z0-9._ -]", "_", name).strip(" .")
    return (name or "upload")[:255]


def _cell_value(cell: dict[str, Any] | None) -> Any:
    return cell.get("value") if cell else None


def _finding(row: int, field: str | None, code: str, message: str, severity: str) -> dict[str, Any]:
    return {"row_number": row, "field": field, "code": code, "message": message, "severity": severity}


def _decimal(value: Any, field: str, row_number: int, findings: list[dict[str, Any]]) -> Decimal | None:
    if value is None or not str(value).strip():
        return None
    normalized = str(value).strip()
    if "," in normalized:
        if not re.fullmatch(r"[+-]?\d{1,3}(?:,\d{3})+(?:\.\d+)?", normalized):
            findings.append(_finding(row_number, field, "ambiguous_number", "Use a plain number or unambiguous thousands separators.", "error"))
            return None
        normalized = normalized.replace(",", "")
    try:
        number = Decimal(normalized)
    except InvalidOperation:
        findings.append(_finding(row_number, field, "invalid_number", "Enter a valid numeric value.", "error"))
        return None
    if not number.is_finite() or abs(number) > Decimal("1000000000000000"):
        findings.append(_finding(row_number, field, "number_out_of_range", "The amount is outside the supported range.", "error"))
        return None
    if field in NONNEGATIVE_FIELDS and number < 0:
        findings.append(_finding(row_number, field, "negative_amount", "This financial amount cannot be negative.", "error"))
        return None
    return number


def validate_mapping(record: PortfolioImport, mapping: dict[str, str | None], worksheet: str) -> None:
    if record.status != "awaiting_mapping" or not record.staged_data:
        raise ImportStateError("This import is not awaiting a column mapping.")
    sheet = record.staged_data.get(worksheet)
    if sheet is None:
        raise ImportFileError("Choose a worksheet from the uploaded workbook.")
    headers = sheet["headers"]
    if set(mapping) - set(headers):
        raise ImportFileError("The mapping contains a column not present in the selected worksheet.")
    targets = [target for target in mapping.values() if target is not None]
    if any(target not in CANONICAL_FIELDS for target in targets):
        raise ImportFileError("The mapping contains an unsupported destination field.")
    if len(targets) != len(set(targets)):
        raise ImportFileError("Each destination field can be mapped only once.")
    if "company_name" not in targets:
        raise ImportFileError("Map a source column to Company name before continuing.")

    column_index = {header: index for index, header in enumerate(headers)}
    field_source = {target: source for source, target in mapping.items() if target is not None}
    findings: list[dict[str, Any]] = []
    validated: list[dict[str, Any]] = []
    seen_keys: set[str] = set()
    for row in sheet["rows"]:
        number = row["row_number"]
        values = row["values"]
        canonical: dict[str, Any] = {}
        row_findings: list[dict[str, Any]] = []
        for source, target in mapping.items():
            if target is None:
                continue
            cell = values[column_index[source]] if column_index[source] < len(values) else None
            if cell and cell.get("formula"):
                row_findings.append(_finding(number, target, "formula_cell", "Formula cells are not imported; replace this cell with a reviewed value.", "error"))
            else:
                canonical[target] = _cell_value(cell)

        company_name = str(canonical.get("company_name") or "").strip()
        if not company_name:
            row_findings.append(_finding(number, "company_name", "required_value", "Company name is required.", "error"))
        else:
            canonical["company_name"] = company_name[:255]
        if not str(canonical.get("industry") or "").strip():
            canonical["industry"] = "Unclassified"
            row_findings.append(_finding(number, "industry", "industry_defaulted", "Industry is not mapped; the company will be labelled Unclassified.", "warning"))

        has_financials = any(canonical.get(field) not in (None, "") for field in FINANCIAL_FIELDS)
        has_monetary_values = any(canonical.get(field) not in (None, "") for field in MONETARY_FIELDS)
        currency = str(canonical.get("currency") or "").strip().upper()
        if has_monetary_values and not currency:
            row_findings.append(_finding(number, "currency", "currency_required", "Map an ISO 4217 currency when importing financial or exposure amounts.", "error"))
        elif currency and currency not in ISO_CURRENCIES:
            row_findings.append(_finding(number, "currency", "invalid_currency", "Enter a recognized ISO 4217 currency code.", "error"))
        elif currency:
            canonical["currency"] = currency
        fiscal_year = canonical.get("fiscal_year")
        if has_financials and (fiscal_year is None or not str(fiscal_year).strip()):
            row_findings.append(_finding(number, "fiscal_year", "reporting_year_required", "A reporting year is required when financial values are provided.", "error"))
        elif fiscal_year is not None and str(fiscal_year).strip():
            try:
                year = int(str(fiscal_year).strip())
                if year < 1900 or year > datetime.now(timezone.utc).year + 1:
                    raise ValueError
                canonical["fiscal_year"] = year
            except ValueError:
                row_findings.append(_finding(number, "fiscal_year", "invalid_reporting_year", "Enter a reporting year between 1900 and next year.", "error"))

        for field in NONNEGATIVE_FIELDS | FINANCIAL_FIELDS | {"interest_rate"}:
            if field not in canonical:
                continue
            number_value = _decimal(canonical[field], field, number, row_findings)
            canonical[field] = str(number_value) if number_value is not None else None
            if field == "interest_rate" and number_value is not None and not Decimal(0) <= number_value <= Decimal(100):
                row_findings.append(_finding(number, field, "interest_rate_out_of_range", "Interest rate must be expressed as a percentage from 0 to 100.", "error"))
        if canonical.get("delinquency_days") not in (None, ""):
            try:
                days = int(str(canonical["delinquency_days"]).strip())
                if days < 0 or days > 36500:
                    raise ValueError
                canonical["delinquency_days"] = days
            except ValueError:
                row_findings.append(_finding(number, "delinquency_days", "invalid_delinquency_days", "Days past due must be a whole number from 0 to 36500.", "error"))

        identity = str(canonical.get("external_id") or "").strip()
        if len(identity) > 128:
            row_findings.append(_finding(number, "external_id", "external_id_too_long", "External IDs must be 128 characters or fewer.", "error"))
        if not identity and company_name:
            identity = "name-" + hashlib.sha256(company_name.casefold().encode("utf-8")).hexdigest()[:32]
        canonical["external_id"] = identity
        if identity in seen_keys:
            row_findings.append(_finding(number, "external_id", "duplicate_in_file", "Another row in this file uses the same company identity; it will be skipped.", "warning"))
        seen_keys.add(identity)

        errors = any(item["severity"] == "error" for item in row_findings)
        validated.append({"row_number": number, "canonical": canonical, "valid": not errors})
        findings.extend(row_findings)

    record.worksheet = worksheet
    record.column_mapping = mapping
    record.findings = findings
    record.total_rows = len(validated)
    record.invalid_rows = sum(not row["valid"] for row in validated)
    record.valid_rows = record.total_rows - record.invalid_rows
    record.warning_count = sum(item["severity"] == "warning" for item in findings)
    record.staged_data = {"rows": validated}
    record.status = "awaiting_confirmation"


def public_import(record: PortfolioImport) -> dict[str, Any]:
    return {
        "id": record.id,
        "filename": record.filename,
        "file_format": record.file_format,
        "file_size": record.file_size,
        "status": record.status,
        "worksheet": record.worksheet,
        "column_mapping": record.column_mapping,
        "total_rows": record.total_rows,
        "valid_rows": record.valid_rows,
        "invalid_rows": record.invalid_rows,
        "warning_count": record.warning_count,
        "imported_rows": record.imported_rows,
        "skipped_rows": record.skipped_rows,
        "failed_rows": record.failed_rows,
        "error_summary": record.error_summary,
        "created_at": record.created_at.isoformat() if record.created_at else None,
        "completed_at": record.completed_at.isoformat() if record.completed_at else None,
    }


def preview_import(record: PortfolioImport, *, limit: int, offset: int) -> dict[str, Any]:
    if record.status != "awaiting_confirmation" or not record.staged_data:
        raise ImportStateError("Validate a column mapping before requesting a preview.")
    rows = record.staged_data["rows"]
    return {
        "total": len(rows),
        "items": rows[offset:offset + limit],
        "findings": [item for item in record.findings if offset < item["row_number"] <= offset + limit],
        "valid_rows": record.valid_rows,
        "invalid_rows": record.invalid_rows,
        "warning_count": record.warning_count,
    }


def confirm_import(db: Session, record: PortfolioImport, tenant_id: int) -> dict[str, Any]:
    if record.status == "completed":
        return public_import(record)
    if record.status != "awaiting_confirmation" or not record.staged_data:
        raise ImportStateError("This import is not ready for confirmation.")

    existing = {
        value for (value,) in db.query(Company.external_id)
        .filter(Company.tenant_id == tenant_id).all()
    }
    seen: set[str] = set()
    imported = 0
    skipped = 0
    for row in record.staged_data["rows"]:
        if not row["valid"]:
            continue
        values = row["canonical"]
        external_id = values["external_id"]
        if external_id in existing or external_id in seen:
            skipped += 1
            continue
        seen.add(external_id)
        company = Company(
            tenant_id=tenant_id,
            external_id=external_id,
            name=values["company_name"],
            industry=values.get("industry") or "Unclassified",
            country=(str(values["country"]).strip()[:100] if values.get("country") else None),
            currency=values.get("currency"),
            annual_revenue=float(Decimal(values["annual_revenue"])) if values.get("annual_revenue") is not None else (
                float(Decimal(values["revenue"])) if values.get("revenue") is not None else None
            ),
            is_demo=False,
            data_source="customer_csv_import",
        )
        db.add(company)
        db.flush()
        financial_values = {field: values.get(field) for field in FINANCIAL_FIELDS}
        if any(value is not None for value in financial_values.values()):
            db.add(CompanyFinancials(
                tenant_id=tenant_id,
                company_id=company.id,
                fiscal_year=values["fiscal_year"],
                revenue=float(Decimal(values["revenue"])) if values.get("revenue") is not None else None,
                ebitda=float(Decimal(values["ebitda"])) if values.get("ebitda") is not None else None,
                net_income=float(Decimal(values["net_income"])) if values.get("net_income") is not None else None,
                total_assets=float(Decimal(values["total_assets"])) if values.get("total_assets") is not None else None,
                total_liabilities=float(Decimal(values["total_liabilities"])) if values.get("total_liabilities") is not None else None,
                current_assets=float(Decimal(values["current_assets"])) if values.get("current_assets") is not None else None,
                current_liabilities=float(Decimal(values["current_liabilities"])) if values.get("current_liabilities") is not None else None,
                debt=float(Decimal(values["debt"])) if values.get("debt") is not None else None,
                cash=float(Decimal(values["cash"])) if values.get("cash") is not None else None,
                operating_cash_flow=float(Decimal(values["operating_cash_flow"])) if values.get("operating_cash_flow") is not None else None,
            ))
        if values.get("requested_loan_amount") is not None or values.get("interest_rate") is not None:
            db.add(CreditProfile(
                tenant_id=tenant_id,
                company_id=company.id,
                requested_loan_amount=float(Decimal(values["requested_loan_amount"])) if values.get("requested_loan_amount") is not None else None,
                interest_rate=float(Decimal(values["interest_rate"])) if values.get("interest_rate") is not None else None,
            ))
        if values.get("exposure") is not None or values.get("outstanding_amount") is not None:
            db.add(PortfolioExposure(
                tenant_id=tenant_id,
                company_id=company.id,
                exposure=float(Decimal(values["exposure"])) if values.get("exposure") is not None else None,
                outstanding_amount=float(Decimal(values["outstanding_amount"])) if values.get("outstanding_amount") is not None else None,
                delinquency_days=values.get("delinquency_days"),
                is_delinquent=bool(values.get("delinquency_days") and values["delinquency_days"] > 0),
            ))
        imported += 1

    record.imported_rows = imported
    record.skipped_rows = skipped
    record.failed_rows = record.invalid_rows
    record.status = "completed_with_errors" if record.invalid_rows else "completed"
    record.completed_at = datetime.now(timezone.utc)
    record.staged_data = None
    db.commit()
    db.refresh(record)
    return public_import(record)


def mapping_suggestions(record: PortfolioImport, worksheet: str | None = None) -> dict[str, Any]:
    if not record.staged_data or record.status != "awaiting_mapping":
        raise ImportStateError("This import is not awaiting a column mapping.")
    selected = worksheet or next(iter(record.staged_data))
    sheet = record.staged_data.get(selected)
    if sheet is None:
        raise ImportFileError("Choose a worksheet from the uploaded workbook.")
    return {
        "worksheets": [name for name in record.staged_data],
        "worksheet": selected,
        "headers": sheet["headers"],
        "sample_rows": [
            {header: _cell_value(row["values"][index] if index < len(row["values"]) else None)
             for index, header in enumerate(sheet["headers"])}
            for row in sheet["rows"][:3]
        ],
        "suggestions": suggest_mapping(sheet["headers"]),
        "fields": CANONICAL_FIELDS,
    }